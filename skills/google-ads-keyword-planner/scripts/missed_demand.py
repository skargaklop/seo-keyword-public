#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Missed demand analysis for Google Ads campaigns.

Subcommands:
  parse-export Parse Google Ads Search Terms export (CSV/XLSX)
  parse-xlsx   Alias for parse-export (backward compatibility)
  build-query  Build OR-query from slot structure
  merge-slots  Merge batch segmentation results into unified slots
  query-total  Estimate demand via Google Ads Historical Metrics
"""
import argparse
import csv
import itertools
import json
import os
import sys
import urllib.request
import urllib.error
import urllib.parse
import re
from collections import OrderedDict


# --- Stop words that may need + prefix in broad query expressions ---
STOP_WORDS = frozenset([
    "на", "в", "к", "за", "с", "по", "из", "от", "до", "для",
    "без", "при", "под", "над", "между", "через", "об", "перед",
])

# Characters forbidden in slot variants (OR-syntax operators and leading modifiers)
# Hyphen inside words is OK (б/у, санкт-петербург)
FORBIDDEN_CHARS_RE = re.compile(r'[|()"]')
LEADING_OPERATOR_RE = re.compile(r'^[+\-!]')
# Punctuation to strip from slot variants (dots, commas, semicolons, etc.)
# Keeps hyphens (б/у, санкт-петербург), slashes, and !
PUNCTUATION_RE = re.compile(r'[.,;:?…·•]')


def _normalize_header(header):
    s = str(header or "").strip().lower()
    s = re.sub(r"\s+", " ", s)
    s = s.replace("_", " ").replace("-", " ")
    return s


def _is_search_term_header(header):
    h = _normalize_header(header)
    return (
        "search term" in h
        or "search query" in h
        or "поисковый запрос" in h
        or "поисковая фраза" in h
    )


def _is_campaign_header(header):
    h = _normalize_header(header)
    return "campaign" in h or "кампан" in h


def _is_ad_group_header(header):
    h = _normalize_header(header)
    return ("ad group" in h) or ("group name" in h) or ("групп" in h)


def _is_status_header(header):
    h = _normalize_header(header)
    return "status" in h or "статус" in h


def _to_rows_from_csv(file_path):
    with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        reader = csv.reader(f, dialect)
        rows = [list(r) for r in reader]
    return rows


def _to_rows_from_xlsx(file_path):
    import openpyxl

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([("" if c is None else str(c)) for c in row])
    wb.close()
    return rows


def _extract_groups_from_table(rows, group_filter=None):
    header_idx = None
    indexes = {}

    for i, row in enumerate(rows[:50]):
        for j, cell in enumerate(row):
            if _is_search_term_header(cell):
                header_idx = i
                indexes["search_term"] = j
                break
        if header_idx is not None:
            break

    if header_idx is None:
        sys.exit("Error: can't find search term column in export")

    header = rows[header_idx]
    for j, cell in enumerate(header):
        if "campaign" not in indexes and _is_campaign_header(cell):
            indexes["campaign"] = j
        if "ad_group" not in indexes and _is_ad_group_header(cell):
            indexes["ad_group"] = j
        if "status" not in indexes and _is_status_header(cell):
            indexes["status"] = j

    groups = OrderedDict()
    auto_idx = 1
    for row in rows[header_idx + 1 :]:
        if len(row) <= indexes["search_term"]:
            continue
        search_term = str(row[indexes["search_term"]]).strip()
        if not search_term:
            continue

        campaign = ""
        ad_group = ""
        status = ""
        if "campaign" in indexes and len(row) > indexes["campaign"]:
            campaign = str(row[indexes["campaign"]]).strip()
        if "ad_group" in indexes and len(row) > indexes["ad_group"]:
            ad_group = str(row[indexes["ad_group"]]).strip()
        if "status" in indexes and len(row) > indexes["status"]:
            status = str(row[indexes["status"]]).strip()

        gid = ad_group or f"group-{auto_idx}"
        if gid == f"group-{auto_idx}":
            auto_idx += 1
        key = f"{campaign}::{gid}"

        if key not in groups:
            groups[key] = {
                "id": gid,
                "name": ad_group or gid,
                "campaign": campaign,
                "status": status,
                "group_minus": "",
                "phrases_ordered": OrderedDict(),
            }

        g = groups[key]
        g["phrases_ordered"][search_term] = None
        if not g["status"] and status:
            g["status"] = status

    result_groups = []
    for _, g in groups.items():
        item = {
            "id": g["id"],
            "name": g["name"],
            "campaign": g["campaign"],
            "status": g["status"],
            "group_minus": g["group_minus"],
            "phrases": list(g["phrases_ordered"].keys()),
        }
        if group_filter and group_filter not in (item["id"], item["name"]):
            continue
        result_groups.append(item)

    return result_groups


def cmd_parse_export(args):
    file_path = args.file
    lower = file_path.lower()
    if lower.endswith(".csv"):
        rows = _to_rows_from_csv(file_path)
    elif lower.endswith(".xlsx"):
        rows = _to_rows_from_xlsx(file_path)
    else:
        sys.exit("Error: file must be .csv or .xlsx")

    groups = _extract_groups_from_table(rows, group_filter=args.group)
    result = {
        "campaign_minus": "",
        "campaign_minus_length": 0,
        "groups": groups,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


def cmd_parse_xlsx(args):
    """Backward-compatible alias."""
    cmd_parse_export(args)


def sanitize_variant(variant, slot_name):
    """Sanitize a slot variant. Returns (cleaned, warnings)."""
    warnings = []
    cleaned = variant

    # Strip punctuation (муз. центр -> муз центр)
    if PUNCTUATION_RE.search(cleaned):
        cleaned = PUNCTUATION_RE.sub("", cleaned)
        cleaned = re.sub(r'\s+', ' ', cleaned).strip()

    # Remove forbidden OR-syntax characters (replace with space to avoid merging tokens)
    if FORBIDDEN_CHARS_RE.search(cleaned):
        warnings.append(
            f"slot '{slot_name}': вариант '{variant}' содержит OR-синтаксис, деградирован в многословный"
        )
        cleaned = FORBIDDEN_CHARS_RE.sub(" ", cleaned)
        # Collapse multiple spaces
        cleaned = re.sub(r'\s+', ' ', cleaned).strip()

    # Remove leading operators (+ - !) from each token
    tokens = cleaned.split()
    new_tokens = []
    for t in tokens:
        if LEADING_OPERATOR_RE.match(t):
            stripped = LEADING_OPERATOR_RE.sub("", t)
            if stripped:
                warnings.append(
                    f"slot '{slot_name}': убран ведущий оператор из '{t}'"
                )
                new_tokens.append(stripped)
        else:
            new_tokens.append(t)
    cleaned = " ".join(new_tokens)

    return cleaned.strip(), warnings


def add_stop_word_plus(variant):
    """Add + prefix to stop words in a variant string."""
    tokens = variant.split()
    result = []
    for t in tokens:
        # Don't touch tokens already starting with operator
        if t[0] in ("+", "-", "!"):
            result.append(t)
        elif t.lower() in STOP_WORDS:
            result.append(f"+{t}")
        else:
            result.append(t)
    return " ".join(result)


def normalize_token(token):
    """Normalize a token for comparison: lowercase, strip punctuation."""
    t = token.lower().strip()
    t = PUNCTUATION_RE.sub("", t)
    t = LEADING_OPERATOR_RE.sub("", t)
    t = FORBIDDEN_CHARS_RE.sub("", t)
    return t.strip()


class SlotMerger:
    """Merges slot segmentation results from multiple LLM batches.

    Does NOT perform LLM segmentation itself — only merges
    already-segmented results and provides coverage analysis.
    """

    SLOT_NAMES = ["objects", "actions", "modifiers", "additional"]
    ADDITIONAL_PATTERN_RE = re.compile(
        r'\b(' + '|'.join(STOP_WORDS) + r')\s+(\S+)',
        re.IGNORECASE,
    )

    def __init__(self, phrases):
        self._phrases = list(phrases)
        self._slots = {
            name: OrderedDict() for name in self.SLOT_NAMES
        }

    def merge_batch(self, batch_index, phrase_indexes, slots):
        """Merge slots from one batch into accumulated state.

        Args:
            batch_index: sequential index of this batch
            phrase_indexes: indices into self._phrases covered by batch
            slots: dict with keys from SLOT_NAMES, values are variant lists

        Raises:
            ValueError: if phrase_indexes are out of range
        """
        n = len(self._phrases)
        for pi in phrase_indexes:
            if not isinstance(pi, int) or pi < 0 or pi >= n:
                raise ValueError(
                    f"phrase_index {pi} вне диапазона [0..{n - 1}]"
                )
        for slot_name in self.SLOT_NAMES:
            variants = slots.get(slot_name, [])
            if not isinstance(variants, list):
                continue
            for v in variants:
                v_str = str(v).strip()
                if not v_str:
                    continue
                cleaned, _ = sanitize_variant(v_str, slot_name)
                if not cleaned:
                    continue
                canon = re.sub(r'\s+', ' ', cleaned.lower()).strip()
                slot_dict = self._slots[slot_name]
                if canon not in slot_dict:
                    slot_dict[canon] = {
                        "display": cleaned,
                        "batch_indexes": set(),
                    }
                slot_dict[canon]["batch_indexes"].add(batch_index)

    def _remove_subsets(self):
        """Remove subset variants within each slot (heuristic).

        If normalized tokens of A are a strict subset of tokens of B,
        B is redundant in OR syntax (A already covers wider queries).
        Comparison uses normalized tokens (lower, no punctuation).
        """
        for slot_name in self.SLOT_NAMES:
            slot_dict = self._slots[slot_name]
            keys = list(slot_dict.keys())
            token_sets = {}
            for k in keys:
                tokens = set(
                    normalize_token(t)
                    for t in k.split()
                    if normalize_token(t)
                )
                token_sets[k] = tokens
            to_remove = set()
            for i, k_short in enumerate(keys):
                for j, k_long in enumerate(keys):
                    if i == j or k_long in to_remove:
                        continue
                    ts_short = token_sets[k_short]
                    ts_long = token_sets[k_long]
                    if ts_short and ts_short < ts_long:
                        to_remove.add(k_long)
            for k in to_remove:
                del slot_dict[k]

    def _compute_coverage(self):
        """Compute coverage of original phrases by merged slot variants."""
        all_variant_tokens = set()
        for slot_name in self.SLOT_NAMES:
            for canon in self._slots[slot_name]:
                for t in canon.split():
                    nt = normalize_token(t)
                    if nt:
                        all_variant_tokens.add(nt)

        uncovered_phrases = []
        all_phrase_tokens = set()
        for idx, phrase in enumerate(self._phrases):
            phrase_tokens = set()
            for t in phrase.split():
                nt = normalize_token(t)
                if nt and nt not in STOP_WORDS:
                    phrase_tokens.add(nt)
                    all_phrase_tokens.add(nt)
            if phrase_tokens and not (phrase_tokens & all_variant_tokens):
                uncovered_phrases.append({"index": idx, "phrase": phrase})

        uncovered_tokens = sorted(
            all_phrase_tokens - all_variant_tokens - STOP_WORDS
        )

        additional_variants_lower = set(
            self._slots.get("additional", OrderedDict()).keys()
        )
        additional_patterns = []
        seen_patterns = set()
        for phrase in self._phrases:
            normalized_phrase = PUNCTUATION_RE.sub("", phrase.lower())
            normalized_phrase = re.sub(r'\s+', ' ', normalized_phrase).strip()
            for m in self.ADDITIONAL_PATTERN_RE.finditer(normalized_phrase):
                pattern = re.sub(r'\s+', ' ', m.group(0)).strip()
                if pattern not in seen_patterns:
                    seen_patterns.add(pattern)
                    if pattern not in additional_variants_lower:
                        additional_patterns.append(pattern)

        return {
            "uncovered_phrases": uncovered_phrases,
            "uncovered_tokens": uncovered_tokens,
            "additional_patterns": additional_patterns,
        }

    def finalize(self, max_query_length=4096):
        """Finalize merged slots: deduplicate, remove subsets, build query."""
        self._remove_subsets()

        slots_pre_trim = {}
        for slot_name in self.SLOT_NAMES:
            slots_pre_trim[slot_name] = [
                entry["display"]
                for entry in self._slots[slot_name].values()
            ]

        coverage = self._compute_coverage()

        debug = {}
        for slot_name in self.SLOT_NAMES:
            debug[slot_name] = {
                entry["display"]: sorted(entry["batch_indexes"])
                for entry in self._slots[slot_name].values()
            }

        slot_order = self.SLOT_NAMES
        warnings = []
        trimmed = []

        sanitized_slots = {}
        for slot_name in slot_order:
            clean_variants = []
            for entry in self._slots[slot_name].values():
                clean_variants.append(add_stop_word_plus(entry["display"]))
            sanitized_slots[slot_name] = clean_variants

        def build_or_string(slot_variants):
            if not slot_variants:
                return ""
            if len(slot_variants) == 1:
                return slot_variants[0]
            return "(" + "|".join(slot_variants) + ")"

        def assemble_query(s_slots):
            parts = []
            for name in slot_order:
                or_str = build_or_string(s_slots[name])
                if or_str:
                    parts.append(or_str)
            return " ".join(parts)

        def est_phrases(s_slots):
            product = 1
            for name in slot_order:
                n = len(s_slots[name])
                if n > 0:
                    product *= n
            return product

        trim_order = ["additional", "modifiers", "actions"]
        while est_phrases(sanitized_slots) > 200:
            trimmed_any = False
            for trim_slot in trim_order:
                if len(sanitized_slots[trim_slot]) > 1:
                    removed = sanitized_slots[trim_slot].pop()
                    trimmed.append({
                        "slot": trim_slot, "removed": removed,
                        "reason": "estimated_phrases > 200",
                    })
                    trimmed_any = True
                    break
            if not trimmed_any:
                break

        query = assemble_query(sanitized_slots)
        while len(query) > max_query_length:
            trimmed_any = False
            for trim_slot in trim_order:
                if len(sanitized_slots[trim_slot]) > 1:
                    removed = sanitized_slots[trim_slot].pop()
                    trimmed.append({
                        "slot": trim_slot, "removed": removed,
                        "reason": f"query_length > {max_query_length}",
                    })
                    trimmed_any = True
                    query = assemble_query(sanitized_slots)
                    break
            if not trimmed_any:
                warnings.append(
                    f"query_length {len(query)} > {max_query_length}"
                    " после максимальной обрезки"
                )
                break

        slots_post_trim = {}
        for slot_name in slot_order:
            slots_post_trim[slot_name] = list(sanitized_slots[slot_name])

        return {
            "slots_pre_trim": slots_pre_trim,
            "slots_post_trim": slots_post_trim,
            "query": query,
            "estimated_phrases": est_phrases(sanitized_slots),
            "query_length": len(query),
            "trimmed": trimmed,
            "warnings": warnings,
            "coverage": coverage,
            "debug": debug,
        }


def cmd_merge_slots(args):
    """Merge batch segmentation results from stdin JSON."""
    raw = sys.stdin.read()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        sys.exit(f"Error: невалидный JSON на stdin: {e}")

    phrases = data.get("phrases", [])
    batches = data.get("batches", [])

    if not isinstance(phrases, list) or not phrases:
        sys.exit("Error: phrases должен быть непустым списком строк")
    if not isinstance(batches, list) or not batches:
        sys.exit("Error: batches должен быть непустым списком объектов")
    for p in phrases:
        if not isinstance(p, str):
            sys.exit(f"Error: элемент phrases не строка: {type(p).__name__}")

    merger = SlotMerger(phrases)
    for i, batch in enumerate(batches):
        if not isinstance(batch, dict):
            sys.exit(f"Error: batch[{i}] не объект")
        phrase_indexes = batch.get("phrase_indexes", [])
        if not isinstance(phrase_indexes, list):
            sys.exit(f"Error: batch[{i}].phrase_indexes не список")
        slots = batch.get("slots", {})
        if not isinstance(slots, dict):
            sys.exit(f"Error: batch[{i}].slots не объект")
        try:
            merger.merge_batch(i, phrase_indexes, slots)
        except ValueError as e:
            sys.exit(f"Error: batch[{i}]: {e}")

    result = merger.finalize(max_query_length=args.max_query_length)
    print(json.dumps(result, ensure_ascii=False, indent=2))


def cmd_build_query(args):
    """Build OR-query from slots JSON."""
    try:
        slots = json.loads(args.slots_json)
    except json.JSONDecodeError as e:
        sys.exit(f"Error: невалидный JSON слотов: {e}")

    max_len = args.max_query_length
    warnings = []
    trimmed = []

    # Slot order: actions, objects, modifiers, additional
    slot_order = ["actions", "objects", "modifiers", "additional"]

    # Sanitize all variants
    sanitized_slots = {}
    for slot_name in slot_order:
        raw_variants = slots.get(slot_name, [])
        if not isinstance(raw_variants, list):
            raw_variants = []
        clean_variants = []
        for v in raw_variants:
            v_str = str(v).strip()
            if not v_str:
                continue
            cleaned, san_warnings = sanitize_variant(v_str, slot_name)
            warnings.extend(san_warnings)
            if cleaned:
                # Apply stop-word plus
                cleaned = add_stop_word_plus(cleaned)
                clean_variants.append(cleaned)
        sanitized_slots[slot_name] = clean_variants

    def build_or_string(slot_variants):
        """Build OR part: (a|b|c) or a."""
        if len(slot_variants) == 0:
            return ""
        if len(slot_variants) == 1:
            return slot_variants[0]
        return "(" + "|".join(slot_variants) + ")"

    def assemble_query(s_slots):
        """Assemble full query from slots dict."""
        parts = []
        for name in slot_order:
            or_str = build_or_string(s_slots[name])
            if or_str:
                parts.append(or_str)
        return " ".join(parts)

    def estimated_phrases(s_slots):
        """Calculate estimated phrase count (product of slot sizes)."""
        product = 1
        for name in slot_order:
            n = len(s_slots[name])
            if n > 0:
                product *= n
        return product

    # Trim priority: additional -> modifiers -> actions (objects untouched)
    trim_order = ["additional", "modifiers", "actions"]

    # Check estimated phrases limit
    while estimated_phrases(sanitized_slots) > 200:
        trimmed_any = False
        for trim_slot in trim_order:
            if len(sanitized_slots[trim_slot]) > 1:
                removed = sanitized_slots[trim_slot].pop()
                trimmed.append({"slot": trim_slot, "removed": removed, "reason": "estimated_phrases > 200"})
                trimmed_any = True
                break
        if not trimmed_any:
            break

    # Check query length limit
    query = assemble_query(sanitized_slots)
    while len(query) > max_len:
        trimmed_any = False
        for trim_slot in trim_order:
            if len(sanitized_slots[trim_slot]) > 1:
                removed = sanitized_slots[trim_slot].pop()
                trimmed.append({"slot": trim_slot, "removed": removed, "reason": f"query_length > {max_len}"})
                trimmed_any = True
                query = assemble_query(sanitized_slots)
                break
        if not trimmed_any:
            warnings.append(
                f"query_length {len(query)} всё ещё > {max_len} после максимальной обрезки"
            )
            break

    result = {
        "query": query,
        "estimated_phrases": estimated_phrases(sanitized_slots),
        "query_length": len(query),
        "trimmed": trimmed,
        "warnings": warnings,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


def _normalize_customer_id(raw):
    return re.sub(r"[\s-]", "", str(raw or ""))


def _refresh_google_access_token(args):
    payload = urllib.parse.urlencode(
        {
            "client_id": args.client_id,
            "client_secret": args.client_secret,
            "refresh_token": args.refresh_token,
            "grant_type": "refresh_token",
        }
    ).encode("utf-8")
    req = urllib.request.Request(
        "https://oauth2.googleapis.com/token",
        data=payload,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            body = resp.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        raw = e.read().decode("utf-8", errors="ignore")[:500]
        raise RuntimeError(f"OAuth HTTP {e.code}: {raw}") from e
    except urllib.error.URLError as e:
        raise RuntimeError(f"OAuth URL error: {e.reason}") from e

    obj = json.loads(body)
    token = obj.get("access_token")
    if not token:
        raise RuntimeError(obj.get("error_description") or obj.get("error") or "access_token missing")
    return token


def _google_ads_json_request(args, endpoint, body):
    access_token = _refresh_google_access_token(args)
    url = f"https://googleads.googleapis.com/{args.api_version}/{endpoint}"
    headers = {
        "Authorization": f"Bearer {access_token}",
        "developer-token": args.developer_token,
        "Content-Type": "application/json",
    }
    if args.login_customer_id:
        headers["login-customer-id"] = _normalize_customer_id(args.login_customer_id)

    req = urllib.request.Request(
        url,
        data=json.dumps(body).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raw = e.read().decode("utf-8", errors="ignore")[:800]
        raise RuntimeError(f"Google Ads HTTP {e.code}: {raw}") from e
    except urllib.error.URLError as e:
        raise RuntimeError(f"Google Ads URL error: {e.reason}") from e


def _cleanup_token_piece(text):
    value = str(text or "").strip()
    value = value.strip("\"'()")
    value = LEADING_OPERATOR_RE.sub("", value)
    value = PUNCTUATION_RE.sub("", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def _expand_or_query(phrase, max_keywords):
    raw_parts = [p for p in str(phrase or "").split() if p]
    slot_options = []

    for part in raw_parts:
        if part.startswith("-"):
            continue
        if part.startswith("(") and part.endswith(")") and "|" in part:
            options = []
            for variant in part[1:-1].split("|"):
                cleaned = _cleanup_token_piece(variant)
                if cleaned and not cleaned.startswith("-"):
                    options.append(cleaned)
            if options:
                slot_options.append(options)
        else:
            cleaned = _cleanup_token_piece(part)
            if cleaned and not cleaned.startswith("-"):
                slot_options.append([cleaned])

    if not slot_options:
        one = _cleanup_token_piece(phrase)
        return [one] if one else []

    keywords = []
    for combo in itertools.product(*slot_options):
        keyword = " ".join(combo).strip()
        if keyword:
            keywords.append(keyword)
        if len(keywords) >= max_keywords:
            break

    seen = OrderedDict()
    for kw in keywords:
        seen[kw.lower()] = kw
    return list(seen.values())


def cmd_query_total(args):
    """Estimate demand via Google Ads historical metrics."""
    customer_id = _normalize_customer_id(args.customer_id)
    if not customer_id:
        sys.exit("Error: customer_id is empty")

    keywords = _expand_or_query(args.phrase, args.max_keywords)
    if not keywords:
        sys.exit("Error: no valid keywords extracted from query")

    body = {
        "keywords": keywords,
        "language": f"languageConstants/{args.language_id}",
        "includeAdultKeywords": False,
    }
    if args.regions:
        try:
            region_ids = [int(r.strip()) for r in args.regions.split(",") if r.strip()]
        except ValueError:
            sys.exit(f"Error: invalid regions: {args.regions}")
        body["geoTargetConstants"] = [f"geoTargetConstants/{r}" for r in region_ids]

    endpoint = f"customers/{customer_id}:generateKeywordHistoricalMetrics"
    try:
        payload = _google_ads_json_request(args, endpoint, body)
    except RuntimeError as e:
        print(json.dumps({"error": str(e), "query": args.phrase}, ensure_ascii=False))
        sys.exit(1)

    if payload.get("error"):
        print(json.dumps({"error": payload["error"], "query": args.phrase}, ensure_ascii=False))
        sys.exit(1)

    total_count = 0
    for row in payload.get("results", []):
        metrics = row.get("keywordMetrics", {})
        total_count += int(metrics.get("avgMonthlySearches", 0) or 0)

    result = {
        "total_count": total_count,
        "query": args.phrase,
        "keyword_count": len(keywords),
        "truncated": len(keywords) >= args.max_keywords,
    }
    print(json.dumps(result, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description="Missed demand analysis tools")
    sub = parser.add_subparsers(dest="command", required=True)

    # parse-export
    p_parse_export = sub.add_parser("parse-export", help="Parse Google Ads export (CSV/XLSX)")
    p_parse_export.add_argument("file", help="Path to CSV/XLSX export")
    p_parse_export.add_argument("--group", default=None, help="Filter by group ID or group name")

    # parse-xlsx (legacy alias)
    p_parse = sub.add_parser("parse-xlsx", help="Alias for parse-export")
    p_parse.add_argument("file", help="Path to CSV/XLSX export")
    p_parse.add_argument("--group", default=None, help="Filter by group ID")

    # build-query
    p_build = sub.add_parser("build-query", help="Build OR-query from slots")
    p_build.add_argument("slots_json", help="JSON string with slots")
    p_build.add_argument("--max-query-length", type=int, default=4096, help="Max query length (default: 4096)")

    # merge-slots
    p_merge = sub.add_parser(
        "merge-slots",
        help="Merge batch segmentation results (stdin JSON)",
    )
    p_merge.add_argument(
        "--max-query-length", type=int, default=4096,
        help="Max query length (default: 4096)",
    )

    # query-total (Google Ads Historical Metrics)
    p_query = sub.add_parser("query-total", help="Estimate total demand via Google Ads historical metrics")
    p_query.add_argument("--developer-token", required=True, help="Google Ads developer token")
    p_query.add_argument("--client-id", required=True, help="OAuth client ID")
    p_query.add_argument("--client-secret", required=True, help="OAuth client secret")
    p_query.add_argument("--refresh-token", required=True, help="OAuth refresh token")
    p_query.add_argument("--customer-id", required=True, help="Google Ads customer ID")
    p_query.add_argument("--login-customer-id", default=None, help="Manager account ID (for MCC)")
    p_query.add_argument("--api-version", default="v22", help="Google Ads API version (default: v22)")
    p_query.add_argument("--phrase", required=True, help="Search phrase with operators")
    p_query.add_argument("--regions", default=None, help="Region IDs comma-separated")
    p_query.add_argument("--language-id", default="1000", help="Language constant ID (default: 1000)")
    p_query.add_argument("--max-keywords", type=int, default=500, help="Keyword expansion cap (default: 500)")

    args = parser.parse_args()
    if args.command == "parse-export":
        cmd_parse_export(args)
    elif args.command == "parse-xlsx":
        cmd_parse_xlsx(args)
    elif args.command == "build-query":
        cmd_build_query(args)
    elif args.command == "merge-slots":
        cmd_merge_slots(args)
    elif args.command == "query-total":
        cmd_query_total(args)


if __name__ == "__main__":
    main()
